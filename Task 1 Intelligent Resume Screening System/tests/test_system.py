"""
Automated Test Suite for Intelligent Resume Screening System
Tests parser, extractor, and matcher modules for correctness and robustness.
"""

import unittest
import io
import docx
from src.parser import extract_text, extract_text_from_docx, _clean_extracted_text
from src.extractor import (
    extract_skills,
    extract_email,
    extract_phone,
    extract_candidate_name,
    categorize_skills,
    SKILL_CATEGORIES,
)
from src.matcher import (
    calculate_tfidf_similarity,
    analyze_skill_gap,
    rank_candidates,
)


class TestResumeScreeningSystem(unittest.TestCase):

    def setUp(self):
        self.sample_resume_text = """
        Dr. Alex Rivera
        Email: alex.rivera@example.com | Phone: +1-555-432-8765
        New York, NY

        Professional Summary:
        Experienced Senior Machine Learning Engineer with 6+ years specializing in NLP and Computer Vision.

        Technical Skills:
        - Languages: Python, SQL, C++, Bash
        - Frameworks: Scikit-Learn, PyTorch, TensorFlow, Pandas, NumPy
        - Cloud & DevOps: Docker, Kubernetes, AWS, Git, CI/CD
        - Databases: PostgreSQL, MongoDB, Redis

        Experience:
        Lead ML Engineer | TechCorp (2020 - Present)
        - Deployed scalable transformer models using Hugging Face and FastAPI.
        - Automated model training pipelines with Docker and Git CI/CD.
        """

        self.sample_jd_text = """
        Looking for a Senior Machine Learning Engineer.
        Required Skills:
        - Strong background in Python, Machine Learning, Deep Learning, and NLP.
        - Experience with PyTorch, Scikit-Learn, and Pandas.
        - Containerization using Docker and AWS deployments.
        - Relational database knowledge (PostgreSQL, SQL).
        - Familiarity with Git and CI/CD pipelines.
        """

    def test_text_cleaning(self):
        raw = "Hello \x00 world!  \r\n\r\n\r\n This is a   test. \xa0 "
        cleaned = _clean_extracted_text(raw)
        self.assertNotIn("\x00", cleaned)
        self.assertNotIn("\xa0", cleaned)
        self.assertIn("Hello world!", cleaned)

    def test_docx_parser(self):
        # Create an in-memory docx document
        doc = docx.Document()
        doc.add_heading("Candidate Resume", level=1)
        doc.add_paragraph("Alice Johnson")
        doc.add_paragraph("Email: alice@test.com")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Skills:"
        table.rows[0].cells[1].text = "Python, Docker, Kubernetes"

        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)

        extracted = extract_text_from_docx(stream)
        self.assertIn("Alice Johnson", extracted)
        self.assertIn("Python", extracted)
        self.assertIn("Docker", extracted)
        self.assertIn("Kubernetes", extracted)

    def test_skill_extraction_positive_matches(self):
        skills = extract_skills(self.sample_resume_text)
        self.assertIn("Python", skills)
        self.assertIn("Machine Learning", skills)
        self.assertIn("Scikit-Learn", skills)
        self.assertIn("PyTorch", skills)
        self.assertIn("Docker", skills)
        self.assertIn("AWS", skills)
        self.assertIn("PostgreSQL", skills)
        self.assertIn("Git", skills)

    def test_skill_boundary_protection(self):
        # Text containing words that shouldn't trigger skills like "C" or "R" or "Go" incorrectly
        tricky_text = "The cat went to go read a clear book about car races in the dark."
        skills = extract_skills(tricky_text)
        self.assertNotIn("C", skills)
        self.assertNotIn("R", skills)
        self.assertNotIn("Go / Golang", skills)

    def test_contact_extraction(self):
        email = extract_email(self.sample_resume_text)
        self.assertEqual(email, "alex.rivera@example.com")

        phone = extract_phone(self.sample_resume_text)
        self.assertIsNotNone(phone)
        self.assertIn("555", phone)

    def test_candidate_name_extraction(self):
        name = extract_candidate_name(self.sample_resume_text)
        self.assertIn("Alex Rivera", name)

    def test_skill_gap_analysis(self):
        jd_skills = ["Python", "Machine Learning", "Docker", "Go / Golang"]
        resume_skills = ["Python", "Machine Learning", "Kubernetes"]

        gap = analyze_skill_gap(jd_skills, resume_skills)
        self.assertIn("Python", gap["matched_skills"])
        self.assertIn("Machine Learning", gap["matched_skills"])
        self.assertIn("Docker", gap["missing_skills"])
        self.assertIn("Go / Golang", gap["missing_skills"])
        self.assertIn("Kubernetes", gap["additional_skills"])
        self.assertEqual(gap["skill_match_score"], 50.0)

    def test_tfidf_cosine_similarity(self):
        scores = calculate_tfidf_similarity(
            self.sample_jd_text,
            [self.sample_resume_text, "Unrelated text about culinary arts and baking cakes."]
        )
        self.assertEqual(len(scores), 2)
        # ML resume should have significantly higher similarity than culinary text
        self.assertGreater(scores[0], scores[1])
        # Scores should be in valid percentage range [0.0, 100.0]
        self.assertGreaterEqual(scores[0], 0.0)
        self.assertLessEqual(scores[0], 100.0)
        self.assertGreaterEqual(scores[1], 0.0)
        self.assertLessEqual(scores[1], 100.0)

    def test_rank_candidates(self):
        candidates = [
            {
                "name": "Alex Rivera",
                "filename": "alex.pdf",
                "text": self.sample_resume_text,
                "email": "alex@test.com",
                "phone": "+1-555-123-4567",
            },
            {
                "name": "Culinary Chef",
                "filename": "chef.docx",
                "text": "Executive pastry chef with 10 years experience baking artisan breads and cakes.",
                "email": "chef@test.com",
                "phone": "+1-555-987-6543",
            }
        ]

        df = rank_candidates(self.sample_jd_text, candidates)
        self.assertEqual(len(df), 2)
        # Alex Rivera should be ranked #1
        self.assertEqual(df.iloc[0]["Rank"], 1)
        self.assertEqual(df.iloc[0]["Candidate Name"], "Alex Rivera")
        self.assertGreater(df.iloc[0]["Match Score (%)"], df.iloc[1]["Match Score (%)"])


if __name__ == "__main__":
    unittest.main()
